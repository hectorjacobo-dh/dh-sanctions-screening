"""Text extraction from evidence documents.

Deliberately conservative. The single most dangerous failure mode here is a
document that yields *no* text — a scanned licence, a photograph of a registry
page — being treated as though it corroborated something. Extraction therefore
always reports a status, and "nothing readable" is a first-class outcome that
the cross-checker treats as providing zero support.

No OCR: there is no Tesseract in this environment, so image-only documents are
reported as unreadable rather than silently contributing nothing.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_EXTRACT_CHARS = 400_000

IMAGE_SUFFIXES = frozenset({".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff"})

_XML_TEXT = re.compile(rb"<w:t[^>]*>([^<]*)</w:t>")
_WS = re.compile(r"[ \t ]+")


@dataclass(frozen=True)
class Extraction:
    text: str
    status: str        # "ok" | "empty" | "unsupported" | "error"
    detail: str
    char_count: int

    @property
    def readable(self) -> bool:
        return self.status == "ok" and bool(self.text.strip())


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS.sub(" ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_EXTRACT_CHARS]


def _from_pdf(payload: bytes) -> Extraction:
    try:
        from pypdf import PdfReader
    except ImportError:
        return Extraction("", "unsupported", "pypdf is not installed", 0)
    try:
        reader = PdfReader(io.BytesIO(payload))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                return Extraction("", "error", "the PDF is password protected", 0)
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception as exc:  # noqa: BLE001
                logger.debug("page extraction failed: %s", exc)
        text = _tidy("\n".join(parts))
        if not text:
            return Extraction(
                "",
                "empty",
                f"the PDF has {len(reader.pages)} page(s) but no extractable text — "
                "it is most likely a scan or photograph, which cannot be read "
                "without OCR",
                0,
            )
        return Extraction(text, "ok", f"{len(reader.pages)} page(s)", len(text))
    except Exception as exc:  # noqa: BLE001
        return Extraction("", "error", f"{type(exc).__name__}: {exc}"[:200], 0)


def _from_docx(payload: bytes) -> Extraction:
    """Read a .docx without a dedicated dependency: it is a zip of XML."""
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = [
                n
                for n in archive.namelist()
                if n.startswith("word/") and n.endswith(".xml") and "document" in n
            ]
            if not names:
                return Extraction("", "empty", "no document body found", 0)
            chunks: list[str] = []
            for name in names:
                for match in _XML_TEXT.finditer(archive.read(name)):
                    chunks.append(match.group(1).decode("utf-8", errors="replace"))
        text = _tidy(" ".join(chunks))
        if not text:
            return Extraction("", "empty", "the document contains no text", 0)
        return Extraction(text, "ok", "word document", len(text))
    except Exception as exc:  # noqa: BLE001
        return Extraction("", "error", f"{type(exc).__name__}: {exc}"[:200], 0)


def _from_xlsx(payload: bytes) -> Extraction:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        chunks: list[str] = []
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(c) for c in row if c is not None]
                    if values:
                        chunks.append(" | ".join(values))
        finally:
            workbook.close()
        text = _tidy("\n".join(chunks))
        if not text:
            return Extraction("", "empty", "the workbook contains no values", 0)
        return Extraction(text, "ok", "spreadsheet", len(text))
    except Exception as exc:  # noqa: BLE001
        return Extraction("", "error", f"{type(exc).__name__}: {exc}"[:200], 0)


def _from_eml(payload: bytes) -> Extraction:
    try:
        from email import message_from_bytes
        from email.policy import default

        message = message_from_bytes(payload, policy=default)
        header_lines = [
            f"{field}: {message.get(field)}"
            for field in ("From", "To", "Cc", "Subject", "Date")
            if message.get(field)
        ]
        bodies: list[str] = []
        if message.is_multipart():
            for part in message.walk():
                if part.get_content_type() == "text/plain":
                    bodies.append(part.get_content())
        else:
            bodies.append(message.get_content())
        text = _tidy("\n".join(header_lines + bodies))
        if not text:
            return Extraction("", "empty", "the email has no plain-text body", 0)
        return Extraction(text, "ok", "email message", len(text))
    except Exception as exc:  # noqa: BLE001
        return Extraction("", "error", f"{type(exc).__name__}: {exc}"[:200], 0)


def _from_text(payload: bytes) -> Extraction:
    text = _tidy(payload.decode("utf-8", errors="replace"))
    if not text:
        return Extraction("", "empty", "the file is empty", 0)
    return Extraction(text, "ok", "plain text", len(text))


def _from_csv(payload: bytes) -> Extraction:
    try:
        raw = payload.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(raw)))
        text = _tidy("\n".join(" | ".join(r) for r in rows if any(r)))
        if not text:
            return Extraction("", "empty", "the file is empty", 0)
        return Extraction(text, "ok", f"{len(rows)} row(s)", len(text))
    except Exception:  # noqa: BLE001
        return _from_text(payload)


def extract(payload: bytes, filename: str) -> Extraction:
    """Pull readable text out of an evidence document."""
    suffix = Path(filename or "").suffix.lower()

    if suffix in IMAGE_SUFFIXES:
        return Extraction(
            "",
            "unsupported",
            "images cannot be read without OCR, which is not installed. This "
            "document provides no automatic corroboration and must be reviewed "
            "by eye.",
            0,
        )
    if suffix == ".pdf":
        return _from_pdf(payload)
    if suffix in {".docx", ".pptx"}:
        return _from_docx(payload)
    if suffix in {".xlsx", ".xlsm"}:
        return _from_xlsx(payload)
    if suffix in {".eml"}:
        return _from_eml(payload)
    if suffix in {".csv"}:
        return _from_csv(payload)
    if suffix in {".txt", ".md", ".rtf"}:
        return _from_text(payload)
    if suffix in {".doc", ".xls", ".ppt", ".msg"}:
        return Extraction(
            "",
            "unsupported",
            f"legacy '{suffix}' files use a proprietary binary format that cannot "
            "be read here. Re-save as the modern equivalent (.docx/.xlsx) or as "
            "PDF to enable automatic cross-checking.",
            0,
        )
    return Extraction("", "unsupported", f"'{suffix}' is not a readable type", 0)
