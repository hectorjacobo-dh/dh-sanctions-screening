"""Batch screening service.

Handles the full run: parse upload -> build blocking index over the active
watchlist snapshot -> score -> persist alerts -> stamp the audit chain.

Column mapping is deliberately permissive. The specification calls for
``Entity_Name / UBO_Name / Country / Tax_ID``, while the live Delivery Hero
export is ``GRID / Type / Account Country / Account Status / Created Date /
Legal Name``. Both are accepted, along with common variants, so regional
managers do not have to reshape their export before uploading.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import logging
import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from itertools import chain

from sqlalchemy import select
from sqlalchemy.orm import Session

from .audit import record_event
from .blocking import BlockingIndex
from .config import get_settings
from .geography import resolve_region
from .matching import band_for, score_against_all_names, score_pair
from .models import MatchResult, ScreeningRun, WatchlistTarget, utcnow
from .ingest.runner import active_refresh_ids, active_snapshot
from .normalize import normalize_name

logger = logging.getLogger(__name__)

# Accepted header spellings, lowercased and stripped of separators.
_COLUMN_MAP: dict[str, tuple[str, ...]] = {
    "grid": ("grid", "partnerid", "vendorid", "accountid", "id"),
    "entity_name": (
        "legalname", "entityname", "partnerlegalname", "partnername",
        "companyname", "name", "restaurantname", "vendorname",
    ),
    "ubo_name": ("uboname", "ubo", "beneficialowner", "ultimatebeneficialowner", "ownername"),
    "country": ("accountcountry", "country", "operatingcountry", "market", "countryname"),
    "tax_id": ("taxid", "vatnumber", "taxnumber", "vat", "registrationnumber", "companynumber"),
    "account_status": ("accountstatus", "status"),
    "created_date": ("createddate", "createdat", "onboardingdate", "signupdate"),
    "entity_type": ("type", "accounttype", "entitytype"),
    "region": ("region", "reportingregion"),
}

MAX_UPLOAD_BYTES = 512 * 1024 * 1024


def _canonical_header(name: str) -> str:
    return "".join(ch for ch in str(name).strip().lower() if ch.isalnum())


def _resolve_columns(header: list[str]) -> dict[str, int]:
    lookup = {_canonical_header(col): i for i, col in enumerate(header)}
    resolved: dict[str, int] = {}
    for field_name, candidates in _COLUMN_MAP.items():
        for candidate in candidates:
            if candidate in lookup:
                resolved[field_name] = lookup[candidate]
                break
    return resolved


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        # The Delivery Hero export is semicolon-delimited; prefer whichever
        # separator actually dominates the header line.
        first_line = sample.splitlines()[0] if sample else ""
        return max(",;\t|", key=first_line.count)


# ---------------------------------------------------------------------------
# Reading: CSV and Excel
# ---------------------------------------------------------------------------

EXCEL_SUFFIXES = (".xlsx", ".xlsm")
LEGACY_EXCEL_SUFFIXES = (".xls",)


def _cell_to_text(value: object) -> str:
    """Render an Excel cell as the text a human would see."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        # Dates arrive as datetimes; keep the date unless a real time is set.
        if value.hour or value.minute or value.second:
            return value.isoformat(sep=" ", timespec="seconds")
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        # Excel stores every number as a float; 4712.0 should read "4712".
        return str(int(value))
    return str(value).strip()


def _load_workbook(payload: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError(
            "Excel support requires openpyxl. Run: pip install -r requirements.txt"
        ) from exc
    return load_workbook(io.BytesIO(payload), read_only=True, data_only=True)


def _choose_excel_sheet(payload: bytes) -> str:
    """Pick the sheet that actually holds the partner records.

    Workbooks routinely lead with a summary or pivot tab — the team's own
    adjudication workbook has the region/source pivot on sheet 1 and the match
    detail on sheet 2. Taking the first sheet reads the wrong table entirely, so
    each sheet's most header-like row is scored and the best one wins, with row
    count as the tiebreak.
    """
    workbook = _load_workbook(payload)
    try:
        names = workbook.sheetnames
        if len(names) <= 1:
            return names[0] if names else ""

        best_name, best_key = names[0], (float("-inf"), 0)
        for name in names:
            sheet = workbook[name]
            prefix: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                prefix.append([_cell_to_text(cell) for cell in row])
                if len(prefix) >= HEADER_SCAN_ROWS:
                    break
            if not prefix:
                continue
            key = (max(_header_score(r) for r in prefix), sheet.max_row or 0)
            if key > best_key:
                best_name, best_key = name, key
        return best_name
    finally:
        workbook.close()


def _iter_excel_rows(payload: bytes) -> Iterator[list[str]]:
    """Stream the best-matching .xlsx/.xlsm sheet row by row.

    read_only keeps memory flat on large exports; openpyxl otherwise builds the
    entire sheet in memory.
    """
    sheet_name = _choose_excel_sheet(payload)
    workbook = _load_workbook(payload)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        for row in sheet.iter_rows(values_only=True):
            yield [_cell_to_text(cell) for cell in row]
    finally:
        workbook.close()


def _iter_csv_rows(payload: bytes) -> tuple[Iterator[list[str]], str]:
    text = payload.decode("utf-8-sig", errors="replace")
    delimiter = _sniff_delimiter(text[:8192])
    return csv.reader(io.StringIO(text), delimiter=delimiter), delimiter


def _iter_raw_rows(payload: bytes, filename: str) -> tuple[Iterator[list[str]], str]:
    """Yield raw rows from either format, plus a delimiter label for reporting."""
    lowered = (filename or "").lower()

    if lowered.endswith(LEGACY_EXCEL_SUFFIXES):
        raise ValueError(
            "Legacy .xls files are not supported. Open the file in Excel and use "
            "Save As -> Excel Workbook (.xlsx), then upload again."
        )

    if lowered.endswith(EXCEL_SUFFIXES):
        return _iter_excel_rows(payload), "xlsx"

    # Sniff the content when the name is unhelpful: .xlsx is a ZIP archive.
    if payload[:2] == b"PK" and b"xl/" in payload[:4096]:
        return _iter_excel_rows(payload), "xlsx"

    return _iter_csv_rows(payload)


# ---------------------------------------------------------------------------
# Header-row detection
# ---------------------------------------------------------------------------

HEADER_SCAN_ROWS = 25


def _header_score(row: list[str]) -> float:
    """How much does this row look like a header rather than data?

    Real exports put banners above the header — the UK Sanctions List opens with
    "Report Date: 16-Jun-2026", and the team's own adjudication workbook uses a
    merged title row. Assuming row 1 is the header silently mis-parses both.
    """
    cells = [c.strip() for c in row]
    filled = [c for c in cells if c]
    if len(filled) < 2:
        return 0.0

    score = 0.0

    # Strongest signal: cells that match a column we actually understand.
    known = sum(
        1
        for c in filled
        if any(_canonical_header(c) in cands for cands in _COLUMN_MAP.values())
    )
    score += known * 10.0

    # Headers are short, wordy labels — not numbers, not dates, not long text.
    numeric = sum(1 for c in filled if c.replace(".", "", 1).replace(",", "").isdigit())
    score += (len(filled) - numeric) * 1.5
    score -= numeric * 2.0

    # Density: a header row is mostly populated across its width.
    if cells:
        score += (len(filled) / len(cells)) * 3.0

    # Distinct labels; data rows repeat values far more often.
    if len(filled) > 1:
        score += (len(set(c.lower() for c in filled)) / len(filled)) * 2.0

    over_long = sum(1 for c in filled if len(c) > 60)
    score -= over_long * 3.0
    return score


def _find_header(rows: list[list[str]]) -> int:
    """Index of the most header-like row within the scanned prefix."""
    best_index, best_score = 0, float("-inf")
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score = _header_score(row)
        if score > best_score:
            best_index, best_score = index, score
    return best_index


# ---------------------------------------------------------------------------
# Content-based name-column detection
# ---------------------------------------------------------------------------

_DATE_LIKE = re.compile(r"^\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")
_CODE_LIKE = re.compile(r"^[A-Z0-9]{4,10}$")
# Identifier shapes: G7, V00012, AB0001, INV-4471. Highly distinct — which is
# exactly why distinctness alone cannot be trusted to find a name column.
_ID_LIKE = re.compile(r"^[A-Za-z]{0,4}[-_]?\d+[A-Za-z]{0,3}$")


def _name_likeness(values: list[str]) -> float:
    """Score how much a column looks like it holds organisation or person names.

    Used when no header matches a known spelling — the column may be called
    "Vendor", "Razón Social", "Nom légal" or nothing at all. Judged on shape:
    names are text, multi-word, long-ish and highly distinct, whereas countries
    and statuses repeat, IDs are short codes, and dates are dates.
    """
    filled = [v.strip() for v in values if v and v.strip()]
    if len(filled) < 3:
        return 0.0

    total = len(filled)
    distinct = len(set(v.lower() for v in filled))

    numeric = sum(1 for v in filled if v.replace(".", "", 1).replace(",", "").isdigit())
    dated = sum(1 for v in filled if _DATE_LIKE.match(v))
    coded = sum(1 for v in filled if _CODE_LIKE.match(v))
    ident = sum(1 for v in filled if _ID_LIKE.match(v))
    lengths = [len(v) for v in filled]
    words = [len(v.split()) for v in filled]

    mean_len = sum(lengths) / total
    mean_words = sum(words) / total
    distinct_ratio = distinct / total

    # Hard gate. A legal name is either multi-word or reasonably long; a column
    # of short single tokens is an identifier, however distinct it looks.
    # Without this, a GRID column wins on uniqueness alone and the engine
    # screens account codes against sanctions lists.
    if mean_words < 1.15 and mean_len < 8:
        return 0.0

    # A column that is mostly identifier-shaped is never the name column.
    if ident / total > 0.6:
        return 0.0

    score = 0.0

    # Names are nearly unique; Country/Status/Type columns are not.
    score += distinct_ratio * 45.0

    # Multi-word text is the signature of a legal name.
    if 1.4 <= mean_words <= 8.0:
        score += 22.0
    elif mean_words < 1.2:
        score -= 12.0

    if 8 <= mean_len <= 70:
        score += 16.0
    elif mean_len < 4:
        score -= 20.0

    # Letters should dominate.
    alpha = sum(1 for v in filled if any(ch.isalpha() for ch in v))
    score += (alpha / total) * 18.0

    # Strong negatives.
    score -= (numeric / total) * 45.0
    score -= (dated / total) * 55.0
    score -= (coded / total) * 30.0

    # Populated columns beat mostly-empty ones.
    score += (total / max(len(values), 1)) * 8.0
    return score


def _detect_name_column(
    header: list[str], sample: list[list[str]]
) -> tuple[int | None, str]:
    """Pick the legal-name column by content. Returns (index, explanation)."""
    if not header:
        return None, ""

    scored: list[tuple[float, int]] = []
    for index in range(len(header)):
        column = [row[index] for row in sample if index < len(row)]
        scored.append((_name_likeness(column), index))

    scored.sort(reverse=True)
    if not scored or scored[0][0] < 25.0:
        return None, ""

    best_score, best_index = scored[0]
    label = (header[best_index] or f"column {best_index + 1}").strip()
    runner_up = scored[1][0] if len(scored) > 1 else 0.0
    return best_index, (
        f"detected by content: {label!r} (score {best_score:.0f} "
        f"vs next-best {runner_up:.0f})"
    )


@dataclass
class PartnerRecord:
    grid: str = ""
    entity_name: str = ""
    ubo_name: str = ""
    country: str = ""
    tax_id: str = ""
    account_status: str = ""
    created_date: str = ""
    entity_type: str = ""
    region: str = ""


@dataclass
class ParsedUpload:
    records: list[PartnerRecord] = field(default_factory=list)
    total_rows: int = 0
    skipped_rows: int = 0
    detected_columns: dict[str, str] = field(default_factory=dict)
    delimiter: str = ","
    # 1-based row where the header was found; > 1 when the file has a banner.
    header_row: int = 1
    # Set when the name column was identified from content rather than header.
    detection_note: str = ""


SAMPLE_ROWS_FOR_DETECTION = 400


def parse_upload(
    payload: bytes, *, filename: str = "", max_rows: int | None = None
) -> ParsedUpload:
    """Parse an uploaded partner file (CSV or Excel) into normalised records.

    Three pieces of tolerance, all driven by what real exports actually look
    like:

    * **Either format.** ``.xlsx``/``.xlsm`` and delimited text, detected from
      the filename or the magic bytes.
    * **The header may not be row 1.** Publications and internal workbooks put
      banners and merged titles above it, so the most header-like row within the
      first 25 is used.
    * **The name column may not be named anything recognisable.** If no header
      matches a known spelling, the column is identified from its *contents* —
      names are text, multi-word, long-ish and near-unique, unlike countries,
      statuses, IDs and dates.

    ``max_rows`` bounds the work for the validation pass inside the upload
    request; parsing 1.3M rows synchronously would hold the connection open.
    """
    row_iter, delimiter = _iter_raw_rows(payload, filename)

    # Buffer enough to find the header and profile the columns beneath it.
    buffer: list[list[str]] = []
    for row in row_iter:
        buffer.append([str(c) if c is not None else "" for c in row])
        if len(buffer) >= HEADER_SCAN_ROWS + SAMPLE_ROWS_FOR_DETECTION:
            break

    if not buffer:
        return ParsedUpload(delimiter=delimiter)

    header_index = _find_header(buffer)
    header = [c.strip() for c in buffer[header_index]]
    body_prefix = buffer[header_index + 1 :]

    columns = _resolve_columns(header)
    detection_note = ""

    if "entity_name" not in columns and "ubo_name" not in columns:
        sample = [r for r in body_prefix if any(c.strip() for c in r)]
        index, detection_note = _detect_name_column(header, sample)
        if index is None:
            preview = ", ".join(h for h in header[:15] if h) or "(no header labels)"
            raise ValueError(
                "Could not identify a column holding partner names. Expected a "
                "header such as: Legal Name, Entity_Name, Partner Legal Name or "
                f"UBO_Name. Header row found at row {header_index + 1}: {preview}"
            )
        columns["entity_name"] = index

    def cell(row: list[str], key: str) -> str:
        index = columns.get(key)
        if index is None or index >= len(row):
            return ""
        return row[index].strip().strip('"')

    parsed = ParsedUpload(
        delimiter=delimiter,
        detected_columns={
            k: (header[v] if v < len(header) and header[v] else f"column {v + 1}")
            for k, v in columns.items()
        },
        header_row=header_index + 1,
        detection_note=detection_note,
    )

    for row in chain(body_prefix, row_iter):
        if not row or not any(c.strip() for c in row):
            continue
        parsed.total_rows += 1

        # Checked before the skip paths below: rows with no screenable name
        # must still count toward the cap, otherwise a file whose leading rows
        # are all blank-named is parsed in full during validation.
        if max_rows is not None and parsed.total_rows > max_rows:
            parsed.total_rows -= 1
            break

        record = PartnerRecord(
            grid=cell(row, "grid"),
            entity_name=cell(row, "entity_name"),
            ubo_name=cell(row, "ubo_name"),
            country=cell(row, "country"),
            tax_id=cell(row, "tax_id"),
            account_status=cell(row, "account_status"),
            created_date=cell(row, "created_date"),
            entity_type=cell(row, "entity_type"),
            region=cell(row, "region"),
        )
        # A row with no screenable name contributes nothing but noise.
        if not record.entity_name and not record.ubo_name:
            parsed.skipped_rows += 1
            continue
        parsed.records.append(record)

    return parsed


@dataclass(slots=True)
class _Target:
    id: int
    source: str
    watchlist_id: str
    name: str
    entity_type: str
    subject_type: str
    program: str
    remarks: str
    geography: str
    dob: str
    aliases: str
    # Structured identity as published. ``geography`` is a registered address;
    # ``nationality`` is a legal assertion about the party. They are compared
    # against different things and must not be conflated.
    nationality: str = ""
    place_of_birth: str = ""
    gender: str = ""
    position: str = ""
    identifiers: str = ""
    listing_reason: str = ""
    linked_to: str = ""
    listed_on: str = ""


def _load_targets(db: Session) -> tuple[dict[int, _Target], BlockingIndex]:
    """Materialise the active watchlist snapshot and index it."""
    settings = get_settings()
    refresh_ids = active_refresh_ids(db)
    if not refresh_ids:
        return {}, BlockingIndex()

    index = BlockingIndex(
        max_token_document_frequency=settings.blocking_max_token_document_frequency,
        max_candidates=settings.max_candidates_per_name,
    )
    targets: dict[int, _Target] = {}

    stmt = (
        select(WatchlistTarget)
        .where(WatchlistTarget.refresh_id.in_(refresh_ids))
        .execution_options(yield_per=5000)
    )
    for row in db.execute(stmt).scalars():
        targets[row.id] = _Target(
            id=row.id,
            source=row.source,
            watchlist_id=row.watchlist_id,
            name=row.target_name,
            entity_type=row.entity_type,
            subject_type=row.subject_type,
            program=row.sanctions_program,
            remarks=row.remarks,
            geography=row.geography,
            dob=row.dob,
            aliases=row.aliases or "",
            nationality=row.nationality or "",
            place_of_birth=row.place_of_birth or "",
            gender=row.gender or "",
            position=row.position or "",
            identifiers=row.identifiers or "",
            listing_reason=row.listing_reason or "",
            linked_to=row.linked_to or "",
            listed_on=row.listed_on or "",
        )
        index.add(row.id, row.target_name)
        # Index aliases against the same target so alternate spellings recall.
        if row.aliases:
            for alias in row.aliases.split(" / "):
                if alias.strip():
                    index.add(row.id, alias)

    index.finalise(len(targets))
    return targets, index


def run_screening(
    db: Session,
    run_id: int,
    payload: bytes,
    *,
    actor: str,
) -> None:
    """Execute a screening run end to end. Intended for a background task."""
    settings = get_settings()
    run = db.get(ScreeningRun, run_id)
    if run is None:
        return

    try:
        run.status = "running"
        db.commit()

        parsed = parse_upload(payload, filename=run.uploaded_filename)
        run.total_rows = parsed.total_rows

        targets, index = _load_targets(db)
        if not targets:
            raise ValueError(
                "No active watchlist snapshot. Refresh the watchlists before screening."
            )

        # Deduplicate on (normalised name, screened field, country). Scoring is
        # identical only when all three agree — country must be part of the key
        # because the geographic control is evaluated per partner country. Two
        # identically-named partners in different countries are genuinely
        # different screening questions and must not share a verdict.
        groups: dict[tuple[str, str, str], list[tuple[PartnerRecord, str, str]]] = {}
        for record in parsed.records:
            for value, kind in ((record.entity_name, "Entity"), (record.ubo_name, "UBO")):
                if not value:
                    continue
                normalized = normalize_name(value)
                if not normalized:
                    continue
                key = (normalized, kind, record.country.strip().upper())
                groups.setdefault(key, []).append((record, value, kind))

        run.distinct_names = len(groups)
        db.commit()

        allowed_subjects = {s.strip() for s in (run.subject_types or "").split(",") if s.strip()}

        total_groups = len(groups)
        alerts: list[MatchResult] = []
        matches_written = 0
        processed = 0

        for (normalized, _kind, _country), members in groups.items():
            processed += 1

            candidate_ids = index.candidates(normalized)
            if candidate_ids:
                # Every member of a group shares the same country by construction.
                group_country = members[0][0].country
                # Score against the ORIGINAL spelling, not the normalised form:
                # normalisation removes the legal suffix, which the acronym
                # bridge needs ("United Aircraft Corporation" -> UAC, not UA).
                # Deterministic: members preserve file order.
                representative_name = members[0][1]

                scored: list[tuple[_Target, object, str]] = []
                for target_id in candidate_ids:
                    target = targets.get(target_id)
                    if target is None:
                        continue
                    # Scope control (e.g. skip vessels/aircraft). Recorded on the
                    # run so the report states what was actually screened.
                    if allowed_subjects and target.subject_type not in allowed_subjects:
                        continue
                    # Score the aliases too. The blocking index recalls a
                    # designation by any published spelling, so scoring only
                    # the primary discarded the very name that matched.
                    breakdown, matched_name = score_against_all_names(
                        representative_name,
                        target.name,
                        target.aliases,
                        group_country,
                        target.geography,
                        geo_penalty=run.geo_penalty,
                    )
                    if breakdown.base_score >= run.threshold:
                        scored.append((target, breakdown, matched_name))

                if scored:
                    # Highest final score first; stable tiebreak for auditability.
                    scored.sort(key=lambda row: (-row[1].final_score, row[0].id))
                    for record, original_value, screened_field in members:
                        region = record.region or resolve_region(record.country)
                        for target, breakdown, matched_name in scored:
                            alerts.append(
                                MatchResult(
                                    run_id=run.id,
                                    grid=record.grid[:60],
                                    region=region,
                                    account_country=record.country[:120],
                                    account_status=record.account_status[:60],
                                    created_date=record.created_date[:40],
                                    partner_legal_name=record.entity_name,
                                    ubo_name=record.ubo_name,
                                    tax_id=record.tax_id[:120],
                                    screened_field=screened_field,
                                    target_name=target.name,
                                    entity_type=target.entity_type[:80],
                                    subject_type=target.subject_type,
                                    sanctions_program=target.program,
                                    remarks=target.remarks,
                                    source=target.source,
                                    geography=target.geography,
                                    dob=target.dob[:200],
                                    watchlist_id=target.watchlist_id[:120],
                                    # Copied, not joined at read time: the
                                    # designation as it stood when the officer
                                    # adjudicated must stay legible after the
                                    # regulator amends or delists the party.
                                    nationality=target.nationality,
                                    place_of_birth=target.place_of_birth,
                                    gender=target.gender[:40],
                                    position=target.position,
                                    identifiers=target.identifiers,
                                    listing_reason=target.listing_reason,
                                    linked_to=target.linked_to,
                                    listed_on=target.listed_on[:120],
                                    matched_alias=(
                                        "" if matched_name == target.name
                                        else matched_name[:300]
                                    ),
                                    # Copied from the run onto every alert so
                                    # the worklist filters without a join, and
                                    # an exported row still says who it is
                                    # about after the case is deleted.
                                    counterparty_group=run.counterparty_group,
                                    brand=run.brand,
                                    score=breakdown.final_score,
                                    base_score=breakdown.base_score,
                                    geo_adjustment=breakdown.geo_adjustment,
                                    risk_band=str(
                                        band_for(
                                            breakdown, run.threshold, run.strong_threshold
                                        )
                                    ),
                                    score_breakdown=breakdown.as_dict(),
                                    reason_codes=",".join(breakdown.reasons),
                                    screened_at_utc=utcnow(),
                                )
                            )

            if len(alerts) >= 2000:
                db.bulk_save_objects(alerts)
                db.commit()
                matches_written += len(alerts)
                alerts = []

            if processed % 2000 == 0 or processed == total_groups:
                run.progress = round(processed / max(total_groups, 1) * 100, 2)
                db.commit()

        if alerts:
            db.bulk_save_objects(alerts)
            db.commit()
            matches_written += len(alerts)

        run.total_matches = matches_written
        run.status = "completed"
        run.progress = 100.0
        run.completed_at = utcnow()
        run.watchlist_snapshot = active_snapshot(db)
        db.commit()

        record_event(
            db,
            "SCREENING_RUN_COMPLETED",
            actor=actor,
            entity_ref=run.run_ref,
            payload={
                "input_sha256": run.input_sha256,
                "filename": run.uploaded_filename,
                "total_rows": run.total_rows,
                "distinct_names": run.distinct_names,
                "total_matches": matches_written,
                "threshold": run.threshold,
                "geo_penalty": run.geo_penalty,
                "subject_types": run.subject_types,
                "watchlist_snapshot": run.watchlist_snapshot,
                "detected_columns": parsed.detected_columns,
                "header_row": parsed.header_row,
                "name_column_detection": parsed.detection_note or "matched by header",
            },
        )

    except Exception as exc:  # noqa: BLE001
        db.rollback()
        run = db.get(ScreeningRun, run_id)
        if run is not None:
            run.status = "failed"
            run.error = f"{type(exc).__name__}: {exc}"[:2000]
            run.completed_at = utcnow()
            db.commit()
            record_event(
                db,
                "SCREENING_RUN_FAILED",
                actor=actor,
                entity_ref=run.run_ref,
                payload={"error": run.error},
            )
        logger.exception("Screening run %s failed", run_id)


def sha256_of(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()
