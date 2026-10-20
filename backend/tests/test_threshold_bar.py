"""The threshold bar on a finished run, and the exports that follow it.

A screening of tens of thousands of names takes the better part of an hour, so
lowering the bar must not mean running it again. Matches are stored from a floor
below the threshold, the bar chooses how many are shown, and the exports show
what is on screen.
"""

from __future__ import annotations

import io

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models import Base, MatchResult, ScreeningRun


@pytest.fixture()
def db():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, future=True)()
    yield session
    session.close()


def _run(db, threshold=95, floor=80) -> ScreeningRun:
    run = ScreeningRun(
        run_ref="r" * 36, uploaded_filename="partners.xlsx", uploaded_by="t", input_sha256="x",
        threshold=threshold, strong_threshold=95, score_floor=floor, geo_penalty=0,
        subject_types="Individual,Entity,Unknown", status="completed",
    )
    db.add(run)
    db.flush()
    return run


def _alert(db, run, score, partner="ACME", target="ACME GROUP", **extra) -> MatchResult:
    fields = dict(
        run_id=run.id, grid=f"G{score}", partner_legal_name=partner, target_name=target,
        score=score, base_score=score, geo_adjustment=0, risk_band="Potential Match",
        score_breakdown={}, reason_codes="", source="OFAC SDN", watchlist_id=str(score),
    )
    fields.update(extra)
    row = MatchResult(**fields)
    db.add(row)
    return row


def test_the_stored_floor_is_below_the_threshold() -> None:
    """Screening at 95 keeps everything from 80, so the bar can be lowered."""
    from app.screening import STORAGE_FLOOR

    assert STORAGE_FLOOR == 80
    assert min(95, STORAGE_FLOOR) == 80
    # A run screened below the floor simply stores from its own threshold.
    assert min(70, STORAGE_FLOOR) == 70


def test_the_report_defaults_to_the_threshold_the_run_was_screened_at(db) -> None:
    from app.report import iter_report_rows

    run = _run(db)
    for score in (100, 96, 90, 85, 80):
        _alert(db, run, score)
    db.commit()

    assert len(list(iter_report_rows(db, run))) == 2          # 100 and 96
    assert len(list(iter_report_rows(db, run, 85))) == 4      # lowering the bar reveals more
    assert len(list(iter_report_rows(db, run, 80))) == 5


def test_the_exports_show_what_is_on_screen(db) -> None:
    """CSV and Excel are built from the same rows, so they cannot disagree."""
    from openpyxl import load_workbook

    from app.report import REPORT_COLUMNS, generate_csv, generate_xlsx

    run = _run(db)
    for score in (100, 90, 85):
        _alert(db, run, score)
    db.commit()

    csv_rows = generate_csv(db, run, 85).strip().split("\r\n")
    sheet = load_workbook(io.BytesIO(generate_xlsx(db, run, 85))).active
    assert len(csv_rows) - 1 == sheet.max_row - 1 == 3
    assert [c.value for c in sheet[1]] == REPORT_COLUMNS


def test_the_report_carries_the_full_published_record(db) -> None:
    """So a case can be dismissed without searching again."""
    from app.report import REPORT_COLUMNS, iter_report_rows

    run = _run(db)
    _alert(
        db, run, 100, dob="10 Apr 1981", nationality="Venezuela", place_of_birth="Caracas",
        gender="Male", identifiers="Passport A123 (Venezuela)", listing_reason="Grounds text",
        linked_to="HIZBALLAH", aliases="ABU X / Y Z", geography="Venezuela",
        reason_codes="GEO_CONTRADICTION_DEPRIORITISED",
    )
    db.commit()
    row = dict(zip(REPORT_COLUMNS, next(iter_report_rows(db, run))))
    for column, expected in (
        ("DOB", "10 Apr 1981"), ("Nationality", "Venezuela"), ("Place_Of_Birth", "Caracas"),
        ("Gender", "Male"), ("Identity_Documents", "Passport A123 (Venezuela)"),
        ("Listing_Reason", "Grounds text"), ("Linked_To", "HIZBALLAH"),
        ("Aliases", "ABU X / Y Z"), ("Country_Check", "differs"), ("Name_Match", "Exact"),
    ):
        assert row[column] == expected, column


def test_a_name_starting_with_a_formula_character_cannot_run_in_excel(db) -> None:
    """Partner names come from an uploaded file. A legal name beginning "=" would
    execute as a formula when an officer opened the report."""
    from openpyxl import load_workbook

    from app.report import REPORT_COLUMNS, generate_csv, generate_xlsx

    run = _run(db)
    _alert(db, run, 100, partner='=HYPERLINK("http://evil.example","x")')
    _alert(db, run, 99, partner="@SUM(1+1)")
    db.commit()

    text = generate_csv(db, run)
    assert '"\'=HYPERLINK' in text and '"\'@SUM' in text

    sheet = load_workbook(io.BytesIO(generate_xlsx(db, run))).active
    column = REPORT_COLUMNS.index("Partner Legal Name") + 1
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=column)
        assert cell.data_type == "s", "a text cell must never be stored as a formula"


def test_the_matches_endpoint_filters_on_the_name_score(db) -> None:
    from app.routers.runs import list_matches

    run = _run(db)
    for score in (100, 96, 90, 85, 80):
        _alert(db, run, score)
    db.commit()

    def shown(min_score):
        return list_matches(
            run.run_ref, limit=100, offset=0, region=None, source=None, risk_band=None,
            decision=None, min_score=min_score, db=db, _=None,
        ).total

    assert shown(None) == 2      # the threshold it was screened at
    assert shown(90) == 3
    assert shown(80) == 5


def test_the_report_routes_are_registered() -> None:
    from app.main import app

    paths = {getattr(route, "path", "") for route in app.routes}
    assert "/api/runs/{run_ref}/report.xlsx" in paths
    assert "/api/runs/{run_ref}/report.csv" in paths
